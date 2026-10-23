"""Fill the SPAQ quantity columns in a downloaded HCM stock template.

The Bulk Stock Upload dialog hands out an .xlsx whose 'Stock Data' sheet has one
row per selected destination facility and one column per product. Only the
product quantity columns are editable; the template's own ReadMe sheet warns
that row 2 holds hidden product-variant IDs and must not be touched, and that
rows and columns must not be added, deleted or rearranged. This script edits
nothing but the quantity cells.

Usage:
    python scripts/fill-stock-template.py --spaq1 10 --spaq2 20
    python scripts/fill-stock-template.py --source path/to/template.xlsx --out filled.xlsx

With no --source it picks the most recently modified Stock_Template_*.xlsx from
the user's Downloads folder. Chrome appends " (1)", " (2)" and so on when a file
of the same name already exists, so choosing by name is unreliable - newest wins.

Requires openpyxl (pip install --user openpyxl).
"""

import argparse
import glob
import os
import shutil
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is not installed. Run: python -m pip install --user openpyxl")

SHEET = "Stock Data"
HEADER_ROW = 1
# Row 2 is the hidden product-variant ID row. The first editable data row is 3.
FIRST_DATA_ROW = 3


def newest_download():
    pattern = os.path.join(os.path.expanduser("~"), "Downloads", "Stock_Template_*.xlsx")
    matches = glob.glob(pattern)
    if not matches:
        sys.exit("No Stock_Template_*.xlsx found in Downloads. Run the download flow first.")
    return max(matches, key=os.path.getmtime)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", help="template to fill (default: newest in ~/Downloads)")
    parser.add_argument("--out", help="where to write the filled copy (default: alongside source)")
    parser.add_argument("--spaq1", type=int, default=200, help="quantity for the SPAQ 1 column")
    parser.add_argument("--spaq2", type=int, default=200, help="quantity for the SPAQ 2 column")
    args = parser.parse_args()

    if args.spaq1 < 0 or args.spaq2 < 0:
        sys.exit("Quantities must be whole numbers of 0 or greater.")

    source = args.source or newest_download()
    out = args.out or os.path.join(os.path.dirname(source), "Stock_Template_filled.xlsx")

    # Copy first and edit the copy, so a failed run never corrupts the download
    # and the original stays available to compare against.
    shutil.copy(source, out)

    workbook = openpyxl.load_workbook(out)
    if SHEET not in workbook.sheetnames:
        sys.exit("Sheet %r not found. Sheets present: %s" % (SHEET, workbook.sheetnames))
    sheet = workbook[SHEET]

    headers = [cell.value for cell in sheet[HEADER_ROW]]
    # Headers carry a space - "SPAQ 1", not "SPAQ1".
    try:
        col_spaq1 = headers.index("SPAQ 1") + 1
        col_spaq2 = headers.index("SPAQ 2") + 1
    except ValueError:
        sys.exit("SPAQ 1 / SPAQ 2 columns not found. Headers are: %s" % headers)

    rows_filled = 0
    for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        # A facility row is identifiable by its destination facility code; blank
        # trailing rows are skipped rather than filled with stray quantities.
        if not sheet.cell(row=row, column=headers.index("To (Facility Code)") + 1).value:
            continue
        sheet.cell(row=row, column=col_spaq1).value = args.spaq1
        sheet.cell(row=row, column=col_spaq2).value = args.spaq2
        rows_filled += 1

    if rows_filled == 0:
        sys.exit("No facility rows found to fill - the template looks empty.")

    workbook.save(out)
    print("source      %s" % source)
    print("filled      %s" % out)
    print("SPAQ 1      column %d = %d" % (col_spaq1, args.spaq1))
    print("SPAQ 2      column %d = %d" % (col_spaq2, args.spaq2))
    print("rows filled %d" % rows_filled)


if __name__ == "__main__":
    main()
